"""Sensibilidad financiera del contenido local del BESS en la malla V2: componentes de costo B4 (control) y B3.

Hipótesis de costo, con selección explícita (no cambia `costos_curvas`, ni su selección automática por energía,
ni el libro):

    H_B4   componentes del segmento B4, los que usa la malla (todas sus energías superan 100 MWh). Control:
           debe reproducir los VPN vigentes.
    H_B3   componentes del segmento B3 aplicados a esos mismos tamaños, solo como sensibilidad del costo local.
           No afirma que B3 esté validado para 120–360 MWh ni que sea la conexión técnicamente correcta; tampoco se
           deduce la tensión de conexión del BESS del nivel de tensión tarifario del usuario.

Solo cambian el CAPEX del BESS y su OPEX fijo (fracción vigente de `costos`). Se conservan las trayectorias, los
ahorros operativos, las exportaciones, las emisiones, los arranques, el SOC, el CAPEX y OPEX del PV, las referencias
red + térmica y PV sin BESS, el horizonte, la tasa, la TRM, el valor residual cero y las convenciones de reposición.
No se recalcula ningún despacho: estas hipótesis de inversión no cambian las decisiones de las reglas.

VPN con el motor existente (`carbono.vpn_financiero_y_ajustado` sobre `motor_multianual`), leyendo las mismas
anualidades que `termica_carbono_vpn.evaluar_vpn`. Sin reposiciones, la comprobación independiente es

    ΔVPN = −ΔCAPEX_BESS − Σ_y ΔOPEX_BESS/(1 + r)^y,

igual para el VPN financiero y el ajustado.

Uso, desde la raíz: `python sensibilidad_capex_bess.py --salida resultados_carbono_v2/sensibilidad_capex_bess_20260927`.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path.cwd()))
import carbono as C
import costos
import costos_curvas as CC
import termica_carbono_anual as TA
import termica_carbono_malla as M
import termica_carbono_malla_informe as I
import termica_carbono_vpn as V

MALLA = Path('resultados_carbono_v2/malla_20260926')
LIBRO = Path('datos/CurvasCapexPVBESS.xlsx')
HIPOTESIS = {'H_B4': 'B4', 'H_B3': 'B3'}
COLUMNAS_LIBRO = {'c_sup': 'D', 'c_e_local': 'E', 'c_p_local': 'F', 'fijo': 'G', 'indirectos': 'H', 'epc': 'I',
                  'margen': 'J'}
TOL_VPN_COP = 1.0            # tolerancia absoluta de reproducción (COP; ≈ 0,0003 USD)
TOL_BENEFICIO_COP = 1e-3     # beneficio operativo anual frente a flujos_anuales.csv (lectura de CSV)
CODIGO = ('sensibilidad_capex_bess.py', 'sensibilidad_capex_bess_pruebas.py', 'costos.py', 'costos_curvas.py',
          'carbono.py', 'motor_multianual.py', 'termica_carbono_vpn.py', 'termica_carbono_malla.py',
          'termica_carbono_malla_informe.py', 'termica_carbono_anual.py')


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ------------------------------------------------------------------------------------ parámetros del libro
def parametros_libro(ruta=LIBRO):
    """Componentes de B3 y B4 leídos directamente de Parametros_BESS (filas por segmento), con su celda."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb['Parametros_BESS']
    filas = {ws.cell(f, 1).value: f for f in range(5, 9)}
    par = {}
    for seg in ('B3', 'B4'):
        f = filas[seg]
        par[seg] = {k: float(ws[f'{col}{f}'].value) for k, col in COLUMNAS_LIBRO.items()}
        par[seg]['celdas'] = {k: f'Parametros_BESS!{col}{f}' for k, col in COLUMNAS_LIBRO.items()}
        par[seg]['descripcion'] = ws.cell(f, 2).value
        par[seg]['rango'] = ws.cell(f, 3).value
    sm = wb['Suministro_BESS']
    prima = dict(quiebre=float(sm['A13'].value), pendiente=float(sm['B13'].value))
    wb.close()
    # Deben coincidir con lo que usa el código vigente.
    _, seg_cc, prima_cc = CC._parametros(ruta)
    for seg, i in (('B3', 2), ('B4', 3)):
        assert seg_cc[i]['segmento'] == seg
        for k in COLUMNAS_LIBRO:
            assert par[seg][k] == float(seg_cc[i][k]), (seg, k)
    assert prima == {k: float(v) for k, v in prima_cc.items()}
    return par, prima


def capex_bess_segmento_USD(potencia_MW, energia_MWh, s, prima, factor_sobredim=CC.FACTOR_SOBREDIMENSIONAMIENTO):
    """Misma ecuación que `costos_curvas.capex_bess_USD`, con el segmento elegido de forma explícita."""
    c_rate = potencia_MW/energia_MWh
    factor_prima = 1.0 + prima['pendiente']*max(0.0, c_rate - prima['quiebre'])
    directo = (s['c_sup']*factor_prima*energia_MWh*factor_sobredim*1_000 + s['c_e_local']*energia_MWh*1_000
               + s['c_p_local']*potencia_MW*1_000 + s['fijo'])
    return directo*(1 + s['indirectos'])*(1 + s['epc'])/(1 - s['margen'])


def costos_hipotesis(cfgs, par, prima, trm):
    filas = []
    for cid, cfg in cfgs.items():
        auto = CC._segmento_bess(cfg['E_MWh'], CC._parametros()[1])['segmento']
        for h, seg in HIPOTESIS.items():
            capex = capex_bess_segmento_USD(cfg['P_MW'], cfg['E_MWh'], par[seg], prima)*trm
            filas.append(dict(Config=cid, **cfg, Hipotesis=h, Segmento_costos=seg, Segmento_automatico=auto,
                              CAPEX_BESS_COP=capex, OPEX_BESS_anual_COP=costos.OPEX_BESS_FRACCION_CAPEX*capex,
                              CAPEX_BESS_MUSD=capex/trm/1e6, USD_kWh=capex/trm/(cfg['E_MWh']*1e3),
                              CAPEX_PV_COP=costos.capex_pv(cfg['f_PV']*M.cargar_malla()['pv_referencia_MWp'], trm),
                              OPEX_PV_anual_COP=costos.opex_pv_anual(cfg['f_PV']*M.cargar_malla()['pv_referencia_MWp'], trm)))
    t = pd.DataFrame(filas)
    # Control: con B4, el CAPEX y el OPEX coinciden exactamente con las funciones vigentes.
    for r in t[t.Hipotesis == 'H_B4'].itertuples():
        assert r.Segmento_automatico == 'B4'
        assert abs(r.CAPEX_BESS_COP - costos.capex_bess(r.P_MW, r.E_MWh, trm)) <= 1e-6
        assert abs(r.OPEX_BESS_anual_COP - costos.opex_bess_anual(r.E_MWh, trm, r.P_MW)) <= 1e-6
    return t


# ------------------------------------------------------------------------------------ VPN con el motor existente
_CACHE = {}


def _meta(grupo, n, s, c, a):
    k = ('meta', grupo, n, s, c, a)
    if k not in _CACHE:
        _CACHE[k] = V.leer_meta(MALLA/grupo, n, s, c, a)
    return _CACHE[k]


def _efc(grupo, n, s, c, a, b):
    k = ('efc', grupo, n, s, c, a)
    if k not in _CACHE:
        _CACHE[k] = V.efc_horario(MALLA/grupo, n, s, c, a, b)
    return _CACHE[k]


def evaluar(grupo, grupo_ref, n, s, proyecto, referencia, cfg, capex_bess_COP, opex_bess_COP):
    """Una comparación con CAPEX y OPEX del BESS explícitos; mismas anualidades y mismo motor que la malla."""
    ctx = TA.contexto()
    trm = ctx['TRM']
    _, b, _, pv_MWp = TA.dimensionar(cfg)
    tiene_bess = proyecto in TA.POLITICAS
    con_pv = referencia == 'RED_TERMICA' and proyecto != 'RED_TERMICA'
    capex = (capex_bess_COP if tiene_bess else 0.) + (costos.capex_pv(pv_MWp, trm) if con_pv else 0.)
    opex = (opex_bess_COP if tiene_bess else 0.) + (costos.opex_pv_anual(pv_MWp, trm) if con_pv else 0.)

    def leer(a, esc_pv):
        mp, mr = _meta(grupo, n, s, proyecto, a), _meta(grupo_ref, n, s, referencia, a)
        if proyecto != 'RED_TERMICA':
            assert abs(mp['Escala_PV'] - esc_pv) < 1e-12
        assert mp['Precio_CO2_USD_t'] == mr['Precio_CO2_USD_t'] and mp['Modalidad'] == mr['Modalidad']
        return dict(costo_proy_fin_USD=mp['Costo_financiero_USD'], costo_ref_fin_USD=mr['Costo_financiero_USD'],
                    emis_cobradas_proy_t=mp['Emisiones_cobradas_t'], emis_cobradas_ref_t=mr['Emisiones_cobradas_t'],
                    precio_USD_t=mp['Precio_CO2_USD_t'],
                    efc_h=_efc(grupo, n, s, proyecto, a, b) if tiene_bess else np.zeros(8760))

    v = C.vpn_financiero_y_ajustado(leer, pv_MWp=pv_MWp, P_MW=b['P_MW'] if tiene_bess else 0.,
                                    E_MWh=b['E_MWh'] if tiene_bess else 0., DoD=b['DoD'], RTE=b['RTE'],
                                    capex_COP=capex, opex_anual_COP=opex, trm=trm)
    o = v['salida']
    if o['n_renovaciones']:
        raise AssertionError(f'{grupo} {n} {s} {proyecto}: reposición; su costo exige la convención por segmento')
    beneficios = [float(x) for x in o['tabla'].Beneficio_COP]
    return dict(Grupo=grupo, Escenario_piloto=n, Senal=s, Proyecto=proyecto, Referencia=referencia, CAPEX_COP=capex,
                OPEX_anual_COP=opex, CAPEX_BESS_COP=capex_bess_COP if tiene_bess else 0.,
                OPEX_BESS_anual_COP=opex_bess_COP if tiene_bess else 0.,
                VPN_financiero_COP=v['VPN_financiero_COP'], VPN_ajustado_COP=v['VPN_ajustado_COP'],
                Valor_carbono_evitado_VP_COP=v['Valor_carbono_evitado_VP_COP'],
                VPN_financiero_MUSD=v['VPN_financiero_COP']/trm/1e6, VPN_ajustado_MUSD=v['VPN_ajustado_COP']/trm/1e6,
                TIR=o['TIR'], Payback_anios=o['Payback_anios'], Renovaciones=o['n_renovaciones'],
                Anios_renovacion=json.dumps(o['Anios_renovacion']), EFC_total=o['EFC_total_descargado'],
                Costo_renovaciones_VP_COP=o['Costo_renovaciones_VP_COP'],
                Beneficios_operativos_COP=json.dumps(beneficios))


def vpn_hipotesis(h, costos_t, m):
    cfgs = M.configs(m)
    ct = costos_t[costos_t.Hipotesis == h].set_index('Config')
    filas = []
    for n in m['escenarios_carbono']:
        for s in m['senales']:
            for f in M.niveles_pv(m):
                filas.append(evaluar(M.ref_pv(f), M.REF_RED, n, s, 'PV_SIN_BESS', 'RED_TERMICA', M.cfg_ref_pv(f), 0., 0.))
            for cid, cfg in cfgs.items():
                cb, ob = ct.loc[cid, 'CAPEX_BESS_COP'], ct.loc[cid, 'OPEX_BESS_anual_COP']
                for regla in m['reglas']:
                    filas.append(evaluar(cid, M.REF_RED, n, s, regla, 'RED_TERMICA', cfg, cb, ob))
                    filas.append(evaluar(cid, M.ref_pv(cfg['f_PV']), n, s, regla, 'PV_SIN_BESS', cfg, cb, ob))
    t = pd.DataFrame(filas)
    t.insert(0, 'Hipotesis', h)
    for c in ('f_PV', 'P_MW', 'E_MWh'):
        t[c] = [cfgs[g][c] if g in cfgs else (int(g[-3:])/100 if c == 'f_PV' else 0.) for g in t.Grupo]
    t['PV_MWp'] = t.f_PV*m['pv_referencia_MWp']
    return t


# ------------------------------------------------------------------------------------ verificaciones
def verificar(vpn, costos_t, m):
    base = pd.read_csv(MALLA/'vpn.csv')
    clave = ['Grupo', 'Escenario_piloto', 'Senal', 'Proyecto', 'Referencia']
    ver = {}
    h4 = vpn[vpn.Hipotesis == 'H_B4'].merge(base, on=clave, suffixes=('', '_base'))
    assert len(h4) == len(base) == 600
    difs = {c: float((h4[c] - h4[f'{c}_base']).abs().max()) for c in ('VPN_financiero_COP', 'VPN_ajustado_COP',
                                                                          'Valor_carbono_evitado_VP_COP', 'CAPEX_COP')}
    assert all(d <= TOL_VPN_COP for d in difs.values()), difs
    ver['H_B4_reproduce_vpn_csv'] = dict(filas=len(h4), tolerancia_COP=TOL_VPN_COP, diferencia_max_COP=difs)
    # Beneficios operativos: iguales entre hipótesis y a flujos_anuales.csv (misma anualidad, mismo beneficio).
    fl = pd.read_csv(MALLA/'flujos_anuales.csv')
    ben = fl.pivot_table(index=['Config', 'Escenario_piloto', 'Senal', 'Proyecto', 'Referencia'], columns='Anio',
                         values='Beneficio_financiero_COP')
    max_ben = 0.
    for h in HIPOTESIS:
        for r in vpn[vpn.Hipotesis == h].itertuples():
            b = np.array(json.loads(r.Beneficios_operativos_COP))
            max_ben = max(max_ben, float(np.abs(b - ben.loc[(r.Grupo, r.Escenario_piloto, r.Senal, r.Proyecto,
                                                              r.Referencia)].to_numpy()).max()))
    assert max_ben <= TOL_BENEFICIO_COP, max_ben   # beneficios anuales del orden de 1e10 COP: último bit ≈ 1e-5
    ver['beneficios_operativos_iguales_a_flujos_anuales'] = dict(diferencia_max_COP=max_ben, tolerancia_COP=TOL_BENEFICIO_COP)
    # Carbono, reposiciones y diferencias entre hipótesis.
    a, b = (vpn[vpn.Hipotesis == h].set_index(clave) for h in ('H_B4', 'H_B3'))
    d = b[['VPN_financiero_COP', 'VPN_ajustado_COP', 'Valor_carbono_evitado_VP_COP', 'CAPEX_COP', 'OPEX_anual_COP']] - \
        a[['VPN_financiero_COP', 'VPN_ajustado_COP', 'Valor_carbono_evitado_VP_COP', 'CAPEX_COP', 'OPEX_anual_COP']]
    assert float(d.Valor_carbono_evitado_VP_COP.abs().max()) == 0.
    ver['componente_carbono_igual_entre_hipotesis'] = True
    assert int(vpn.Renovaciones.sum()) == 0 and int(base.Renovaciones.sum()) == 0 and \
        int(fl.Renovaciones_en_anio.sum()) == 0
    ver['reposiciones'] = 'ninguna en los VPN originales ni en las dos hipótesis (EFC < 7.200 y vida calendaria = horizonte)'
    factor = sum(1/(1 + costos.WACC)**y for y in range(1, costos.HORIZONTE_ANIOS + 1))
    analitico = -d.CAPEX_COP - d.OPEX_anual_COP*factor
    err_f = float((d.VPN_financiero_COP - analitico).abs().max())
    err_a = float((d.VPN_ajustado_COP - analitico).abs().max())
    assert err_f <= TOL_VPN_COP and err_a <= TOL_VPN_COP, (err_f, err_a)
    assert float((d.VPN_financiero_COP - d.VPN_ajustado_COP).abs().max()) <= TOL_VPN_COP
    ver['delta_vpn_analitico'] = dict(factor_anualidad=factor, error_max_financiero_COP=err_f,
                                      error_max_ajustado_COP=err_a,
                                      formula='ΔVPN = −ΔCAPEX_BESS − ΔOPEX_BESS·Σ_y 1/(1+r)^y')
    # Aditividad en cada hipótesis.
    max_ad = 0.
    for h in HIPOTESIS:
        x = vpn[vpn.Hipotesis == h]
        pv = x[x.Proyecto == 'PV_SIN_BESS'].set_index(['Escenario_piloto', 'Senal', 'f_PV'])
        for (g, n, s, p), y in x[x.Proyecto.isin(m['reglas'])].groupby(['Grupo', 'Escenario_piloto', 'Senal', 'Proyecto']):
            y = y.set_index('Referencia')
            for col in ('VPN_financiero_COP', 'VPN_ajustado_COP'):
                dif = y.loc['RED_TERMICA', col] - pv.loc[(n, s, y.f_PV.iloc[0]), col] - y.loc['PV_SIN_BESS', col]
                max_ad = max(max_ad, abs(float(dif)))
    assert max_ad <= TOL_VPN_COP, max_ad
    ver['aditividad'] = dict(diferencia_max_COP=max_ad)
    pvs = d.xs('PV_SIN_BESS', level='Proyecto')
    assert float(pvs[['VPN_financiero_COP', 'VPN_ajustado_COP']].abs().max().max()) == 0.
    ver['solo_PV_sin_cambio'] = True
    return ver, d.reset_index()


# ------------------------------------------------------------------------------------ comparaciones y figuras
def comparaciones(vpn, m, salida):
    """Comparaciones A, B y C de la malla, para cada hipótesis, con las funciones del informe de la malla."""
    I.TRM = TA.contexto()['TRM']
    ops = pd.read_csv(MALLA/'operacion_20_anios.csv')
    ops = ops[ops.Anualidades == M.H]
    res = {}
    for h in HIPOTESIS:
        carpeta = Path(salida)/h
        carpeta.mkdir(parents=True, exist_ok=True)
        r = I.resultados_config(vpn[vpn.Hipotesis == h].drop(columns='Hipotesis'), ops, m)
        r.to_csv(carpeta/'resultados_por_configuracion.csv', index=False)
        I.escribir_resumenes(carpeta, r, m, 'COMPLETO')
        res[h] = r
    return res


def tablas_eleccion(salida):
    salida = Path(salida)
    lee = lambda h, f: pd.read_csv(salida/h/f)
    cc = lee('H_B4', 'comparacion_configuraciones.csv').merge(
        lee('H_B3', 'comparacion_configuraciones.csv'), on=['Escenario_piloto', 'Precio_CO2_USD_t', 'Senal'],
        suffixes=('_B4', '_B3'))
    cc['Cambia_ganador'] = cc.Config_1_B4 != cc.Config_1_B3
    cc['Cambia_segunda'] = cc.Config_2_B4 != cc.Config_2_B3
    br = lee('H_B4', 'mejores_por_regla.csv').merge(lee('H_B3', 'mejores_por_regla.csv'),
                                                     on=['Escenario_piloto', 'Precio_CO2_USD_t', 'Senal', 'Regla'],
                                                     suffixes=('_B4', '_B3'))
    br['Cambia_tamano'] = br.Config_B4 != br.Config_B3
    ar = lee('H_B4', 'reglas_tamano_comun.csv').merge(lee('H_B3', 'reglas_tamano_comun.csv'),
                                                      on=['Config', 'Escenario_piloto', 'Senal'], suffixes=('_B4', '_B3'))
    reglas = ('A_INT', 'B_INT', 'D_INT')
    orden = lambda r, h: [k for k in sorted(reglas, key=lambda k: -r[f'{k}_VPN_ajustado_MUSD_{h}'])]
    ar['Orden_B4'] = [','.join(orden(r, 'B4')) for _, r in ar.iterrows()]
    ar['Orden_B3'] = [','.join(orden(r, 'B3')) for _, r in ar.iterrows()]
    ar['Mismo_orden'] = ar.Orden_B4 == ar.Orden_B3
    ar['Cambio_max_diferencia_entre_reglas_MUSD'] = [
        max(abs((r[f'{x}_VPN_ajustado_MUSD_B3'] - r[f'{y}_VPN_ajustado_MUSD_B3'])
                - (r[f'{x}_VPN_ajustado_MUSD_B4'] - r[f'{y}_VPN_ajustado_MUSD_B4'])) for x in reglas for y in reglas)
        for _, r in ar.iterrows()]
    cc.to_csv(salida/'eleccion_configuraciones_B4_B3.csv', index=False)
    br.to_csv(salida/'eleccion_tamano_por_regla_B4_B3.csv', index=False)
    ar.to_csv(salida/'reglas_tamano_fijo_B4_B3.csv', index=False)
    return cc, br, ar


def figuras(salida, res, br, m):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    plt.rcParams.update({'font.size': 8.5, 'axes.spines.top': False, 'axes.spines.right': False, 'axes.grid': True,
                         'grid.color': '#deded9', 'grid.linewidth': 0.6, 'savefig.facecolor': 'white'})
    carpeta = Path(salida)/'figuras'
    carpeta.mkdir(exist_ok=True)
    fmt = plt.FuncFormatter(lambda v, _: I._es(v))
    col = {'H_B4': '#7048c0', 'H_B3': '#c07a00'}
    mk = {'H_B4': 'o', 'H_B3': 'D'}
    ids = list(M.configs(m))
    x = np.arange(len(ids))
    esc = m['escenarios_carbono']
    # 1. Mejor VPN ajustado de cada configuración (mejor regla) con las dos hipótesis; ganador marcado.
    fig, axs = plt.subplots(2, len(esc), figsize=(13, 7.4), sharex=True)
    for i, s in enumerate(m['senales']):
        for j, n in enumerate(esc):
            ax = axs[i, j]
            textos = []
            for h in HIPOTESIS:
                r = res[h]
                g = r[(r.Senal == s) & (r.Escenario_piloto == n)]
                best = g.loc[g.groupby('Config').VPN_ajustado_MUSD.idxmax()].set_index('Config').reindex(ids)
                ax.plot(x + (-0.12 if h == 'H_B4' else 0.12), best.VPN_ajustado_MUSD, mk[h], ms=4.5, color=col[h],
                        mfc=col[h] if h == 'H_B4' else 'white', mew=1.2)
                w = best.VPN_ajustado_MUSD.idxmax()
                ax.plot(ids.index(w) + (-0.12 if h == 'H_B4' else 0.12), best.loc[w, 'VPN_ajustado_MUSD'], 'o', ms=11,
                        mfc='none', mec=col[h], mew=1.4)
                textos.append(f"{h[2:]}: {w} {best.loc[w, 'Regla'].split('_')[0]}")
            ax.set_title(f"{s} — {I.PRECIO[n]} USD/t\nganador {' · '.join(textos)}", fontsize=8.5)
            ax.yaxis.set_major_formatter(fmt)
            ax.set_xticks(x, ids, fontsize=7, rotation=90)
            if j == 0:
                ax.set_ylabel('VPN ajustado, proyecto completo [MUSD]\n(mejor regla de cada configuración)')
    fig.legend(handles=[Line2D([], [], marker=mk[h], ls='', color=col[h], mfc=col[h] if h == 'H_B4' else 'white',
                               label=f'{h}: componentes locales de {HIPOTESIS[h]}') for h in HIPOTESIS]
               + [Line2D([], [], marker='o', ls='', ms=11, mfc='none', mec='#444444', label='mayor VPN de la malla')],
               loc='lower center', ncol=3, frameon=False)
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    fig.savefig(carpeta/'01_eleccion_configuracion_B4_B3.png', dpi=150)
    plt.close(fig)
    # 2. Mejor tamaño por regla: celdas donde cambia entre hipótesis.
    filas = [(r, s) for r in m['reglas'] for s in m['senales']]
    mat = np.zeros((len(filas), len(esc)))
    fig, ax = plt.subplots(figsize=(10, 4.6))
    for a, (r, s) in enumerate(filas):
        for b_, n in enumerate(esc):
            z = br[(br.Regla == r) & (br.Senal == s) & (br.Escenario_piloto == n)].iloc[0]
            mat[a, b_] = float(z.Cambia_tamano)
            ax.text(b_, a, f'B4: {z.Config_B4}\nB3: {z.Config_B3}', ha='center', va='center', fontsize=8,
                    fontweight='bold' if z.Cambia_tamano else 'normal')
    ax.imshow(mat, cmap=matplotlib.colors.ListedColormap(['#f0efea', '#f2d7a6']), vmin=0, vmax=1, aspect='auto')
    ax.set_xticks(range(len(esc)), [f'{I.PRECIO[n]} USD/t' for n in esc])
    ax.set_yticks(range(len(filas)), [f'{r} — {s}' for r, s in filas])
    ax.grid(False)
    ax.set_title('Mejor tamaño encontrado para cada regla con componentes locales B4 y B3 (sombreado: cambia)',
                 fontsize=9)
    fig.tight_layout()
    fig.savefig(carpeta/'02_mejor_tamano_por_regla_B4_B3.png', dpi=150)
    plt.close(fig)
    return sorted(str(p) for p in carpeta.glob('*.png'))


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument('--salida', type=Path, required=True)
    a = p.parse_args(argv)
    salida = a.salida
    if salida.resolve() == MALLA.resolve():
        raise SystemExit('La sensibilidad no escribe en la carpeta de la malla')
    salida.mkdir(parents=True, exist_ok=True)
    antes = {f: sha(MALLA/f) for f in ('vpn.csv', 'flujos_anuales.csv', 'anualidades.csv',
                                        'resultados_por_configuracion.csv', 'operacion_20_anios.csv', 'manifiesto.json')}
    m = M.cargar_malla()
    trm = TA.contexto()['TRM']
    par, prima = parametros_libro()
    pd.DataFrame([dict(Segmento=s, Descripcion=par[s]['descripcion'], Rango_declarado=par[s]['rango'], Componente=k,
                       Celda=par[s]['celdas'][k], Valor=par[s][k],
                       Difiere_entre_B3_y_B4=par['B3'][k] != par['B4'][k]) for s in ('B3', 'B4') for k in COLUMNAS_LIBRO]
                 + [dict(Segmento='ambos', Componente=k, Celda=f'Suministro_BESS!{c}13', Valor=prima[k],
                         Difiere_entre_B3_y_B4=False) for k, c in (('quiebre', 'A'), ('pendiente', 'B'))]
                 ).to_csv(salida/'parametros_segmentos.csv', index=False)
    ct = costos_hipotesis(M.configs(m), par, prima, trm)
    ct.to_csv(salida/'capex_opex_por_configuracion.csv', index=False)
    vpn = pd.concat([vpn_hipotesis(h, ct, m) for h in HIPOTESIS], ignore_index=True)
    ver, d = verificar(vpn, ct, m)
    vpn.drop(columns='Beneficios_operativos_COP').to_csv(salida/'vpn_hipotesis.csv', index=False)
    d['Delta_VPN_MUSD'] = d.VPN_ajustado_COP/trm/1e6
    d.rename(columns={c: f'Delta_{c}' for c in ('VPN_financiero_COP', 'VPN_ajustado_COP', 'Valor_carbono_evitado_VP_COP',
                                                 'CAPEX_COP', 'OPEX_anual_COP')}).to_csv(salida/'diferencias_vpn_B3_menos_B4.csv',
                                                                                          index=False)
    res = comparaciones(vpn, m, salida)
    # Control de las tablas: H_B4 reproduce las tablas vigentes de la malla.
    for f in ('comparacion_configuraciones.csv', 'mejores_por_regla.csv', 'reglas_tamano_comun.csv'):
        x, y = pd.read_csv(salida/'H_B4'/f), pd.read_csv(MALLA/f)
        num = [c for c in y.select_dtypes('number').columns if c in x.columns]
        txt = [c for c in y.select_dtypes(exclude='number').columns if c in x.columns and c != 'Estado_entrega']
        assert x[txt].equals(y[txt]), f
        assert float((x[num] - y[num]).abs().max().max()) <= 1e-6, f
    ver['H_B4_reproduce_tablas_de_la_malla'] = 'comparacion_configuraciones, mejores_por_regla y reglas_tamano_comun (≤ 1e-6 MUSD)'
    cc, br, ar = tablas_eleccion(salida)
    figs = figuras(salida, res, br, m)
    despues = {f: sha(MALLA/f) for f in antes}
    assert antes == despues, 'la sensibilidad modificó resultados de la malla'
    ver['huellas_malla_sin_cambio'] = despues
    ver['resumen'] = dict(cambia_ganador=int(cc.Cambia_ganador.sum()), cambia_tamano_por_regla=int(br.Cambia_tamano.sum()),
                          reglas_mismo_orden=int(ar.Mismo_orden.sum()), comparaciones_reglas=len(ar))
    (salida/'verificaciones.json').write_text(json.dumps(ver, ensure_ascii=False, indent=1, default=str), encoding='utf-8')
    archivos = sorted(p for p in salida.rglob('*') if p.is_file() and p.name != 'manifiesto.json')
    (salida/'manifiesto.json').write_text(json.dumps(dict(
        generado=pd.Timestamp.now().isoformat(), hipotesis=HIPOTESIS, alcance=__doc__.split('\n\n')[1],
        entradas={**{f'malla/{f}': h for f, h in antes.items()}, 'CurvasCapexPVBESS.xlsx': sha(LIBRO)},
        resultados={str(p.relative_to(salida)): sha(p) for p in archivos},
        codigo={f: sha(f) for f in CODIGO if Path(f).exists()},
        comando='python sensibilidad_capex_bess.py --salida ' + str(salida), figuras=figs),
        ensure_ascii=False, indent=1, default=str), encoding='utf-8')
    (salida/'ESTADO.txt').write_text('COMPLETO' + chr(10), encoding='utf-8')
    print(json.dumps(ver['resumen'], ensure_ascii=False), flush=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
