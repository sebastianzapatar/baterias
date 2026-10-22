"""CAPEX de ``CurvasCapexPVBESS.xlsx`` y OPEX porcentual solicitado.

Esta capa sustituye, solo para el análisis multi-sitio, los supuestos planos de
``costos.py``. Lee los parámetros del libro en cada ejecución para que cambiar de
equipo o actualizar el Excel no exija copiar cifras dentro del código.

Convenciones del análisis:

* USD constantes de 2026, escenario medio y TRM promedio de 2025 del proyecto.
* PV propio, alcance Full EPC con 5 km de línea MT.
* BESS de 4 horas, sin sobredimensionamiento adicional (factor 1,00), de modo que
  el caso 10 MW / 40 MWh reproduce exactamente la hoja ``Curva_BESS``.
* OPEX PV = 2,0 % del CAPEX solar por año.
* OPEX BESS = 1,5 % del CAPEX BESS por año.
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

import costos


LIBRO_CAPEX = Path("datos/CurvasCapexPVBESS.xlsx")
ANIO_COSTOS = 2026
OPEX_PV_FRACCION = 0.020
OPEX_BESS_FRACCION = 0.015
FACTOR_SOBREDIMENSIONAMIENTO = 1.0


def _numero(valor, etiqueta):
    if not isinstance(valor, (int, float)):
        raise ValueError(f"{etiqueta}: se esperaba un número y se obtuvo {valor!r}")
    return float(valor)


@lru_cache(maxsize=4)
def _parametros(ruta: str | Path = LIBRO_CAPEX):
    """Lee las celdas fuente del libro; no depende de sus fórmulas calculadas."""
    ruta = Path(ruta)
    if not ruta.exists():
        raise FileNotFoundError(
            f"No se encontró {ruta}. Ejecute el análisis desde la raíz de OperacionBESS."
        )

    wb = load_workbook(ruta, read_only=True, data_only=True)

    pv = wb["Parametros_PV"]
    parametros_pv = {
        "hardware": _numero(pv["C16"].value, "Parametros_PV!C16"),
        "local": _numero(pv["C17"].value, "Parametros_PV!C17"),
        "conexion": _numero(pv["C18"].value, "Parametros_PV!C18"),
        "factor": _numero(pv["C20"].value, "Parametros_PV!C20"),
        "b_hardware": _numero(pv["C33"].value, "Parametros_PV!C33"),
        "b_local": _numero(pv["C34"].value, "Parametros_PV!C34"),
    }

    bess = wb["Parametros_BESS"]
    segmentos = []
    for fila in range(5, 9):
        segmentos.append(
            {
                "segmento": str(bess.cell(fila, 1).value),
                "c_sup": _numero(bess.cell(fila, 4).value, f"Parametros_BESS!D{fila}"),
                "c_e_local": _numero(bess.cell(fila, 5).value, f"Parametros_BESS!E{fila}"),
                "c_p_local": _numero(bess.cell(fila, 6).value, f"Parametros_BESS!F{fila}"),
                "fijo": _numero(bess.cell(fila, 7).value, f"Parametros_BESS!G{fila}"),
                "indirectos": _numero(bess.cell(fila, 8).value, f"Parametros_BESS!H{fila}"),
                "epc": _numero(bess.cell(fila, 9).value, f"Parametros_BESS!I{fila}"),
                "margen": _numero(bess.cell(fila, 10).value, f"Parametros_BESS!J{fila}"),
            }
        )

    suministro = wb["Suministro_BESS"]
    prima = {
        "quiebre": _numero(suministro["A13"].value, "Suministro_BESS!A13"),
        "pendiente": _numero(suministro["B13"].value, "Suministro_BESS!B13"),
    }
    wb.close()
    return parametros_pv, segmentos, prima


def capex_pv_unitario_usd_kwp(pv_MWp: float, ruta=LIBRO_CAPEX) -> float:
    """Curva solar continua por escala, ecuación declarada en ``Parametros_PV``."""
    if pv_MWp <= 0:
        return 0.0
    p, _, _ = _parametros(Path(ruta))
    escala = pv_MWp / 8.0
    return (
        p["hardware"] * escala ** p["b_hardware"]
        + p["local"] * escala ** p["b_local"]
        + p["conexion"]
    ) * p["factor"]


def capex_pv_COP(pv_MWp: float, trm=costos.TRM_COP_USD, ruta=LIBRO_CAPEX) -> float:
    return capex_pv_unitario_usd_kwp(pv_MWp, ruta) * pv_MWp * 1_000 * trm


def _segmento_bess(energia_MWh: float, segmentos):
    if energia_MWh <= 0:
        return None
    if energia_MWh < 1:
        return segmentos[0]
    if energia_MWh < 20:
        return segmentos[1]
    if energia_MWh <= 100:
        return segmentos[2]
    return segmentos[3]


def capex_bess_USD(
    potencia_MW: float,
    energia_MWh: float,
    factor_sobredimensionamiento=FACTOR_SOBREDIMENSIONAMIENTO,
    ruta=LIBRO_CAPEX,
) -> float:
    """Modelo bottom-up por segmento declarado en ``Parametros_BESS``."""
    if potencia_MW < 0 or energia_MWh < 0:
        raise ValueError("La potencia y la energía del BESS no pueden ser negativas.")
    if energia_MWh == 0:
        if potencia_MW != 0:
            raise ValueError("Un BESS sin energía también debe tener potencia cero.")
        return 0.0
    if potencia_MW <= 0:
        raise ValueError("Un BESS con energía debe tener potencia positiva.")

    _, segmentos, prima = _parametros(Path(ruta))
    s = _segmento_bess(energia_MWh, segmentos)
    c_rate = potencia_MW / energia_MWh
    factor_prima = 1.0 + prima["pendiente"] * max(0.0, c_rate - prima["quiebre"])
    energia_instalada_kWh = energia_MWh * factor_sobredimensionamiento * 1_000
    energia_nominal_kWh = energia_MWh * 1_000
    potencia_kW = potencia_MW * 1_000

    costo_directo = (
        s["c_sup"] * factor_prima * energia_instalada_kWh
        + s["c_e_local"] * energia_nominal_kWh
        + s["c_p_local"] * potencia_kW
        + s["fijo"]
    )
    return costo_directo * (1 + s["indirectos"]) * (1 + s["epc"]) / (1 - s["margen"])


def capex_bess_COP(potencia_MW: float, energia_MWh: float, trm=costos.TRM_COP_USD, **kwargs) -> float:
    return capex_bess_USD(potencia_MW, energia_MWh, **kwargs) * trm


def resumen_costos(pv_MWp: float, potencia_MW: float, energia_MWh: float, trm=costos.TRM_COP_USD):
    capex_pv = capex_pv_COP(pv_MWp, trm)
    capex_bess = capex_bess_COP(potencia_MW, energia_MWh, trm)
    return {
        "CAPEX_PV_COP": capex_pv,
        "CAPEX_BESS_COP": capex_bess,
        "CAPEX_total_COP": capex_pv + capex_bess,
        "OPEX_PV_anual_COP": OPEX_PV_FRACCION * capex_pv,
        "OPEX_BESS_anual_COP": OPEX_BESS_FRACCION * capex_bess,
        "OPEX_total_anual_COP": OPEX_PV_FRACCION * capex_pv + OPEX_BESS_FRACCION * capex_bess,
        "CAPEX_PV_USD_kWp": capex_pv_unitario_usd_kwp(pv_MWp),
        "CAPEX_BESS_USD_kWh": (
            capex_bess_USD(potencia_MW, energia_MWh) / (energia_MWh * 1_000)
            if energia_MWh > 0 else 0.0
        ),
    }


# ============================================================
# Bloque de almacenamiento reemplazable (R1) y trayectorias
# ============================================================


def capex_bess_bloque_almacenamiento(
    potencia_MW: float,
    energia_MWh: float,
    factor_sobredimensionamiento=FACTOR_SOBREDIMENSIONAMIENTO,
    ruta=LIBRO_CAPEX,
) -> dict:
    """Costo del bloque de almacenamiento reemplazable, desglosado.

    Es la parte del CAPEX que se vuelve a comprar cuando se agota la envolvente
    de utilización del BESS: los términos que escalan con ENERGÍA, con sus
    multiplicadores.

        R1 = (cSUP · prima(C) · E_inst + cE_local · E_nom)
             × (1 + indirectos) × (1 + EPC) ÷ (1 − margen)

    Quedan FUERA `cP_local · P` y `cF`, que representan equipo y contenido local
    que escala con potencia —transformador, BOP eléctrico, obra civil, SCADA— y
    el costo fijo del proyecto. Ninguno de los dos se agota por ciclado.

    **No es «costo de celdas».** El propio libro documenta que el suministro se
    cotiza por energía y que el PCS viene embebido en el bloque AC, así que R1
    arrastra una parte de equipo de potencia. Es, por tanto, una aproximación
    CONSERVADORA del costo del bloque reemplazable con la estructura de costos
    disponible: sobreestima lo que hay que reponer, no lo subestima.

    Devuelve los sumandos por separado porque cada uno sigue una trayectoria
    tecnológica distinta (ver `indice_trayectoria`).
    """
    if potencia_MW <= 0 or energia_MWh <= 0:
        return {"suministro_USD": 0.0, "local_energia_USD": 0.0,
                "directo_USD": 0.0, "sobrecostos_USD": 0.0, "total_R1_USD": 0.0,
                "multiplicador": 1.0}

    _, segmentos, prima = _parametros(Path(ruta))
    s = _segmento_bess(energia_MWh, segmentos)
    c_rate = potencia_MW / energia_MWh
    factor_prima = 1.0 + prima["pendiente"] * max(0.0, c_rate - prima["quiebre"])

    suministro = (s["c_sup"] * factor_prima
                  * energia_MWh * factor_sobredimensionamiento * 1_000)
    local_energia = s["c_e_local"] * energia_MWh * 1_000
    directo = suministro + local_energia
    mult = (1 + s["indirectos"]) * (1 + s["epc"]) / (1 - s["margen"])

    return {
        "suministro_USD": suministro * mult,
        "local_energia_USD": local_energia * mult,
        "directo_USD": directo,
        "sobrecostos_USD": directo * (mult - 1.0),
        "total_R1_USD": directo * mult,
        "multiplicador": mult,
        "segmento": s["segmento"],
    }


_SERIES_TRAYECTORIA = {"bateria": 2, "solar_utility": 5, "solar_comercial": 8,
                       "local": 11}
_ESCENARIOS_TRAYECTORIA = {"bajo": 0, "medio": 1, "alto": 2}


@lru_cache(maxsize=4)
def _trayectorias(ruta: str | Path = LIBRO_CAPEX):
    """Índices anuales de la hoja `Trayectorias`, en USD constantes.

    Base 2026 = 1,00. La hoja los publica año a año hasta 2050.
    """
    wb = load_workbook(Path(ruta), read_only=True, data_only=True)
    hoja = wb["Trayectorias"]
    tabla = {}
    for fila in hoja.iter_rows(min_row=15, max_row=60, max_col=13, values_only=True):
        if not fila or not isinstance(fila[0], (int, float)):
            continue
        tabla[int(fila[0])] = tuple(
            float(v) if isinstance(v, (int, float)) else float("nan")
            for v in fila[1:13]
        )
    wb.close()
    if not tabla:
        raise ValueError("La hoja Trayectorias no devolvió ningún año.")
    return tabla


# ------------------------------------------------------------------
# Trayectoria del suministro BESS: serie anual publicada por NREL
# ------------------------------------------------------------------
# C3E. La hoja `Trayectorias` no contiene la serie de NREL: la RECONSTRUYE con
# una tasa compuesta constante por tramo. Esa reconstrucción se desvía poco en
# el escenario Alto (+0,1 % en 2035) pero mucho en el Bajo (-13,5 % en 2035 y
# -57,9 % en 2050), porque extrapola una pendiente inicial agresiva mientras la
# serie real de NREL se aplana.
#
# Desde C3E la trayectoria PRODUCTIVA del suministro se lee de este CSV, que
# transcribe la tabla publicada. Las columnas de batería de la hoja quedan como
# LEGACY y ya no alimentan ningún cálculo.
#
# El Excel sigue siendo la fuente de verdad de todo lo demás: parámetros
# colombianizados (cSUP, cE_local, cP_local, cF, indirectos, EPC, margen),
# segmentos B1-B4 y la TRAYECTORIA LOCAL.
TABLA_NREL = Path("datos/NREL_TP-6A40-93281_tabla2.csv")
_COLUMNAS_NREL = {"bajo": "low_index_2026", "medio": "mid_index_2026",
                  "alto": "high_index_2026"}


@lru_cache(maxsize=4)
def _trayectoria_nrel(ruta: str | Path = TABLA_NREL):
    """Índices anuales del suministro BESS, base `ANIO_COSTOS` = 1,00.

    Los valores no se codifican en Python: se leen del CSV, que a su vez
    transcribe la tabla 2 de NREL/TP-6A40-93281.
    """
    ruta = Path(ruta)
    if not ruta.exists():
        raise FileNotFoundError(
            f"No se encontró {ruta}. Es la trayectoria de suministro BESS y "
            "debe estar en datos/, junto a CurvasCapexPVBESS.xlsx."
        )
    tabla = {}
    with ruta.open(encoding="utf-8") as f:
        encabezado = None
        for linea in f:
            if linea.startswith("#") or not linea.strip():
                continue
            campos = [c.strip() for c in linea.rstrip("\n").split(",")]
            if encabezado is None:
                encabezado = campos
                continue
            fila = dict(zip(encabezado, campos))
            tabla[int(fila["year"])] = {
                k: float(fila[col]) for k, col in _COLUMNAS_NREL.items()
            }
    if not tabla:
        raise ValueError(f"{ruta} no contiene ninguna fila de datos.")
    if ANIO_COSTOS not in tabla:
        raise ValueError(f"{ruta} no cubre el año base {ANIO_COSTOS}.")
    return tabla


def indice_trayectoria(anio: int, serie: str = "bateria",
                       escenario: str = "medio", ruta=LIBRO_CAPEX,
                       ruta_nrel=TABLA_NREL) -> float:
    """Índice de costo real del año pedido, con `ANIO_COSTOS` = 1,00.

    `serie` es una de {bateria, solar_utility, solar_comercial, local} y
    `escenario` una de {bajo, medio, alto}.

    **Dos fuentes, sin ambigüedad**: `bateria` sale del CSV de NREL; el resto
    sigue saliendo de la hoja `Trayectorias` del Excel.

    Fuera del rango disponible lanza `ValueError`: no se extrapola en silencio.
    """
    if serie not in _SERIES_TRAYECTORIA:
        raise ValueError(f"Serie desconocida: {serie!r}")
    if escenario not in _ESCENARIOS_TRAYECTORIA:
        raise ValueError(f"Escenario desconocido: {escenario!r}")

    if serie == "bateria":
        tabla = _trayectoria_nrel(Path(ruta_nrel))
        if anio not in tabla:
            raise ValueError(
                f"El año {anio} está fuera de la serie NREL de suministro "
                f"({min(tabla)}–{max(tabla)}). No se extrapola."
            )
        return tabla[anio][escenario]

    tabla = _trayectorias(Path(ruta))
    if anio not in tabla:
        raise ValueError(
            f"El año {anio} está fuera de la tabla de trayectorias "
            f"({min(tabla)}–{max(tabla)}). No se extrapola."
        )
    col = _SERIES_TRAYECTORIA[serie] + _ESCENARIOS_TRAYECTORIA[escenario] - 2
    return tabla[anio][col]


def validar_con_excel():
    """Guardas de regresión contra dos celdas calculadas del libro."""
    pv_100 = capex_pv_unitario_usd_kwp(100)
    bess_40 = capex_bess_USD(10, 40) / 40_000
    if abs(pv_100 - 679.2012146966388) > 1e-6:
        raise AssertionError(f"La curva PV ya no reproduce 100 MWp: {pv_100}")
    if abs(bess_40 - 178.6001993759375) > 1e-6:
        raise AssertionError(f"La curva BESS ya no reproduce 10 MW/40 MWh: {bess_40}")
    return {"PV_100_USD_kWp": pv_100, "BESS_4h_USD_kWh": bess_40}


if __name__ == "__main__":
    print(validar_con_excel())
